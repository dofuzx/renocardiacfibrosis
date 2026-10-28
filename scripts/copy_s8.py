# -*- coding: utf-8 -*-
"""把 S8_external_validation 表从 pre Review 副本完整复制到 Frontiers IPO 副本
（逐单元格复制值/样式/合并/列宽），补齐被正文引用的 S8。"""
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

SRC = r"D:\AIwork\论文投稿\FP\pre Review\FrontPharmacol\Supplementary_Tables.xlsx"
DST = r"D:\AIwork\论文投稿\FP\Frontiers IPO\Supplementary_Tables.xlsx"
SHEET = "S8_external_validation"


def copy_sheet(src_ws, dst_wb):
    if SHEET in dst_wb.sheetnames:
        del dst_wb[SHEET]
    dst_ws = dst_wb.create_sheet(SHEET)
    # 单元格值 + 样式
    for row in src_ws.iter_rows():
        for c in row:
            cell = dst_ws.cell(row=c.row, column=c.column, value=c.value)
            cell.number_format = c.number_format
            cell.font = c.font.copy()
            cell.fill = c.fill.copy()
            cell.border = c.border.copy()
            cell.alignment = c.alignment.copy()
    # 合并单元格
    for mc in src_ws.merged_cells.ranges:
        dst_ws.merge_cells(str(mc))
    # 列宽
    for col, dim in src_ws.column_dimensions.items():
        dst_ws.column_dimensions[col].width = dim.width
    # 行高
    for ridx, dim in src_ws.row_dimensions.items():
        dst_ws.row_dimensions[ridx].height = dim.height
    return dst_ws


def main():
    swb = load_workbook(SRC)
    dwb = load_workbook(DST)
    copy_sheet(swb[SHEET], dwb)
    dwb.save(DST)
    print("copied", SHEET, "->", DST)
    print("sheets now:", dwb.sheetnames)


if __name__ == "__main__":
    main()
