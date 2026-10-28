# -*- coding: utf-8 -*-
# 生成投稿最终版补充表：S7 状态更新 + 新增 S8（GSE145154 外部人心验证）
# 源文件只读；产物写入 pre Review/FrontPharmacol，不改动原 Supplementary_Tables.xlsx
import openpyxl, csv, os
from openpyxl.styles import Font, Alignment, PatternFill

SRC_XLSX = "D:/AIwork/论文投稿/FP/Frontiers IPO/Supplementary_Tables.xlsx"
SRC_CSV  = "D:/AIwork/论文投稿/FP/phase7_humanheart/results_20261024/"
DST      = "D:/AIwork/论文投稿/FP/pre Review/FrontPharmacol/"
os.makedirs(DST, exist_ok=True)

wb = openpyxl.load_workbook(SRC_XLSX)

# ---- 1) S7: GSE145154 由 Reserve 改为 Analysed -----------------------------------
ws7 = wb["S7_data_resources"]
n = 0
for row in ws7.iter_rows(min_row=3):
    rowvals = [str(c.value) if c.value is not None else "" for c in row]
    if "GSE145154" in rowvals:
        for c in row:
            if c.value is not None and str(c.value).strip():
                if c.column == 1:  # Status
                    c.value = "Analysed (external validation)"
                    c.font = Font(bold=True)
                if c.column == 8:  # Notes
                    c.value = ("External human cardiac validation cohort (§2.10, §3.6, Fig 8); "
                               "24 ventricular libraries from six hearts (1 non-failing donor, 2 DCM, 3 ICM); "
                               "110,355 nuclei after QC")
                if c.column == 7:  # n samples
                    c.value = "24"
        n += 1
print("S7 rows updated:", n)

# ---- 2) 新增 S8 ------------------------------------------------------------------
if "S8_external_validation" in wb.sheetnames:
    del wb["S8_external_validation"]
ws = wb.create_sheet("S8_external_validation")

HDR = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="4472C4")
SUB = Font(bold=True, size=11)
blocks = [
    ("Table S8a. CYP1B1 detection by cardiac lineage (GSE145154, n = 110,355 nuclei)",
     "T1_CYP1B1_by_celltype.csv"),
    ("Table S8b. CYP1B1 detection by lineage and condition (per-cent expressing nuclei, %)",
     "T1b_CYP1B1_by_celltype_condition.csv"),
    ("Table S8c. Fibroblast vs each non-fibroblast compartment (Wilcoxon, BH-adjusted)",
     "T2_CYP1B1_fibro_vs_others.csv"),
    ("Table S8d. Disease response of CYP1B1 within cardiac fibroblasts (vs donor nuclei)",
     "T3_CYP1B1_disease_response_in_fibro.csv"),
    ("Table S8e. Fibroblast pseudo-bulk CYP1B1 per ventricular library",
     "T4_fibro_pseudobulk_per_sample.csv"),
    ("Table S8f. Within-fibroblast coupling of CYP1B1 to activation markers (Spearman)",
     "T5_CYP1B1_activation_correlation.csv"),
]
r = 1
for title, fn in blocks:
    ws.cell(r, 1, title).font = SUB
    r += 1
    with open(SRC_CSV + fn, newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    for i, row in enumerate(rows):
        for j, v in enumerate(row, start=1):
            c = ws.cell(r, j, v)
            if i == 0:
                c.font = HDR
                c.fill = FILL
                c.alignment = Alignment(horizontal="center")
        r += 1
    r += 2
for col, w in zip("ABCDEFGHI", [26, 20, 20, 20, 20, 20, 22, 22, 22]):
    ws.column_dimensions[col].width = w
print("S8 rows written:", r)

# ---- 3) README 追加 S8 -------------------------------------------------------------
wsr = wb["README"]
lastopen = wsr.max_row + 2
wsr.cell(lastopen, 1, "Supplementary Table S8 — External human heart snRNA-seq validation "
                      "(GSE145154): CYP1B1 lineage restriction, disease response, pseudo-bulk "
                      "per library and activation coupling (panels a–f).").font = Font(size=10)

out = DST + "Supplementary_Tables.xlsx"
wb.save(out)
print("saved:", out, os.path.getsize(out), "bytes")
